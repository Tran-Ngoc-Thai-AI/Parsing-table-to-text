from pathlib import Path
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from openai import OpenAI


class DocxProcessor:
    """
    Process one DOCX file.

    Flow:
        DOCX
        -> detect tables
        -> send tables to LLM
        -> validate LLM results
        -> replace tables with semantic text
        -> save processed DOCX
    """

    def __init__(self, settings_path: Path = Path("local.settings.json")):
        settings_path = Path(settings_path)

        with open(settings_path, "r", encoding="utf-8") as f:
            settings = json.load(f)["Values"]

        self.azure_openai_api_key = settings["AZURE_OPENAI_API_KEY"]
        self.azure_openai_deployment = settings["AZURE_OPENAI_DEPLOYMENT"]
        self.azure_openai_endpoint = settings["AZURE_OPENAI_ENDPOINT"]

        self.client = OpenAI(
            api_key=self.azure_openai_api_key,
            base_url=self.azure_openai_endpoint,
        )

    # ============================================================
    # 1. Load DOCX
    # ============================================================

    def load_document(self, input_path: Path):
        input_path = Path(input_path)

        if not input_path.exists():
            raise FileNotFoundError(
                f"Input DOCX not found: {input_path}"
            )

        return Document(input_path)

    # ============================================================
    # 2. Detect tables
    # ============================================================

    def detect_tables(self, doc):
        tables_need_llm = []
        tables_skipped = []

        for i, table in enumerate(doc.tables):

            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in table.rows
            ]

            table_text = " ".join(
                cell
                for row in rows
                for cell in row
                if cell
            )

            row_count = len(rows)
            col_count = len(table.columns)

            # Giữ nguyên logic notebook
            if row_count >= 2 and table_text:
                tables_need_llm.append({
                    "table_index": i,
                    "rows": row_count,
                    "cols": col_count,
                    "text": table_text,
                })
            else:
                if row_count < 2:
                    reason = "Less than 2 rows"
                elif not table_text:
                    reason = "Empty table"
                else:
                    reason = "Does not meet LLM criteria"

                tables_skipped.append({
                    "table_index": i,
                    "rows": row_count,
                    "cols": col_count,
                    "reason": reason,
                })

        return tables_need_llm, tables_skipped

    # ============================================================
    # 3. Prepare tables for LLM
    # ============================================================

    def prepare_tables_for_llm(self, doc, tables_need_llm):

        tables_for_llm = []

        for item in tables_need_llm:

            table = doc.tables[item["table_index"]]

            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in table.rows
            ]

            tables_for_llm.append({
                "table_index": item["table_index"],
                "rows": rows,
            })

        return tables_for_llm

    # ============================================================
    # 4. Build prompt
    # ============================================================

    def build_prompt(self, rows):

        table_json = json.dumps(
            rows,
            ensure_ascii=False,
            indent=2,
        )

        prompt = f"""
Bạn là hệ thống chuyển đổi bảng dữ liệu trong tài liệu nghiệp vụ
thành văn bản có cấu trúc để sử dụng làm dữ liệu cho chatbot/RAG.

NHIỆM VỤ

Đọc bảng đầu vào và chuyển toàn bộ thông tin trong bảng thành
structured semantic text.

Mục tiêu là:
- Giữ nguyên thông tin của bảng.
- Làm rõ quan hệ giữa các trường và giá trị.
- Loại bỏ sự phụ thuộc vào cấu trúc bảng khi chatbot tìm kiếm.
- Cho phép người đọc hiểu được nội dung mà không cần nhìn bảng gốc.

QUY TRÌNH XỬ LÝ BẮT BUỘC

Bước 1 — Xác định record

Hãy coi MỖI DÒNG của bảng là một record độc lập.

Không được gộp các giá trị của các dòng khác nhau ở bước này.

Bước 2 — Xác định entity của từng record

Từ các cột trong bảng, xác định các trường tạo nên "đối tượng chính"
của record.

Các trường mô tả định danh hoặc nhận diện đối tượng phải được giữ
nguyên cùng nhau.

Các trường có thể bao gồm:
- tên đối tượng
- mã đối tượng
- tên tiếng Anh
- số hiệu
- hoặc các trường khác có vai trò nhận diện đối tượng.

Không được tự động coi một cột đơn lẻ là entity nếu điều đó làm
mất quan hệ giữa các cột.

Bước 3 — Group record

Chỉ group hai hoặc nhiều record khi chúng mô tả CÙNG MỘT ĐỐI TƯỢNG.

Hai record KHÔNG được coi là cùng đối tượng nếu các giá trị định danh
chính của chúng khác nhau.

Đặc biệt:

KHÔNG được group toàn bộ bảng theo từng cột.

KHÔNG được lấy tất cả giá trị của cột A gom lại với tất cả giá trị
của cột B.

KHÔNG được tạo kết quả kiểu:

Thuộc tính A: A1 / A2
Thuộc tính B: B1 / B2
Thuộc tính C: C1 / C2 / C3

nếu các giá trị này thuộc các record khác nhau.

Bước 4 — Merge thuộc tính

Sau khi đã xác định các record thuộc CÙNG MỘT ĐỐI TƯỢNG,
mới được merge các giá trị khác nhau của cùng một thuộc tính.

Ví dụ:

Record 1:
Entity A | X | Value 1

Record 2:
Entity A | X | Value 2

=> merge thành:

Entity: A
X: Value 1 / Value 2

Nhưng:

Record 1:
Entity A | X | Value 1

Record 2:
Entity B | X | Value 2

=> KHÔNG được merge.

Phải tạo hai block riêng.

Bước 5 — Xuất kết quả

Mỗi entity chỉ có MỘT block.

Trong một block:
- Mỗi thuộc tính chỉ xuất hiện một lần.
- Nếu thuộc tính có nhiều giá trị do nhiều record của cùng entity,
  merge các giá trị bằng "/ ".
- Các thuộc tính khác nhau giữa các entity tuyệt đối không được
  trộn lẫn với nhau.

QUY TẮC QUAN TRỌNG NHẤT

Thứ tự xử lý bắt buộc là:

ROWS
→ IDENTIFY ENTITY
→ GROUP ROWS BY ENTITY
→ MERGE ATTRIBUTES WITHIN EACH ENTITY
→ OUTPUT BLOCKS

Tuyệt đối KHÔNG xử lý theo:

COLUMNS
→ MERGE VALUES
→ OUTPUT

Nếu không chắc hai record có cùng entity hay không,
KHÔNG được merge chúng.

YÊU CẦU VỀ KHẢ NĂNG SEARCH

Kết quả phải giúp chatbot có thể tìm kiếm riêng biệt các thông tin
quan trọng trong bảng.

Do đó:
- Các đối tượng khác nhau phải có thể phân biệt được.
- Các thuộc tính quan trọng phải được thể hiện rõ.
- Các giá trị thuộc cùng một đối tượng phải được liên kết rõ ràng.
- Không làm mất thông tin chỉ vì thông tin đó bị lặp lại giữa
  nhiều dòng.
- Khi nhiều giá trị thuộc cùng một thuộc tính của cùng một đối tượng,
  phải đặt các giá trị đó trên cùng một dòng thuộc tính.

KIỂM TRA TRƯỚC KHI TRẢ KẾT QUẢ

Trước khi trả về kết quả, phải kiểm tra:

- Có bỏ sót thông tin nào trong bảng không?
- Có tạo thêm thông tin không có trong bảng không?
- Các giá trị trên cùng một dòng có còn đúng quan hệ không?
- Các dòng thuộc cùng một đối tượng đã được gộp đúng chưa?
- Trong mỗi block có thuộc tính nào bị lặp lại không?
- Nếu một thuộc tính có nhiều giá trị, các giá trị đó đã được
  merge vào cùng một dòng chưa?

Nếu phát hiện một thuộc tính xuất hiện nhiều lần trong cùng một block,
phải gộp các giá trị của thuộc tính đó thành một dòng duy nhất
trước khi trả kết quả.

BẢNG ĐẦU VÀO:

{table_json}
"""

        return prompt

    # ============================================================
    # 5. Reconstruct tables with LLM
    # ============================================================

    def reconstruct_tables(self, tables_for_llm):

        results = []

        for table_data in tables_for_llm:

            table_index = table_data["table_index"]
            rows = table_data["rows"]

            prompt = self.build_prompt(rows)

            response = self.client.chat.completions.create(
                model=self.azure_openai_deployment,
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "Bạn là chuyên gia chuyển đổi dữ liệu bảng "
                            "nghiệp vụ thành structured semantic text "
                            "cho hệ thống RAG."
                        ),
                    },
                    {
                        "role": "user",
                        "content": prompt,
                    },
                ],
                temperature=0,
            )

            semantic_text = (
                response.choices[0].message.content.strip()
            )

            results.append({
                "table_index": table_index,
                "semantic_text": semantic_text,
            })

            # print()
            # print("=" * 80)
            # print(f"TABLE {table_index}")
            # print("=" * 80)
            # print(semantic_text)

        return results

    # ============================================================
    # 6. Validate LLM results
    # ============================================================

    def validate_results(self, tables_for_llm, results):

        print()
        print(f"Expected tables : {len(tables_for_llm)}")
        print(f"LLM results     : {len(results)}")

        assert len(results) == len(tables_for_llm), (
            f"Missing results: "
            f"expected {len(tables_for_llm)}, "
            f"got {len(results)}"
        )

        for result in results:

            table_index = result["table_index"]
            semantic_text = result["semantic_text"]

            assert semantic_text, (
                f"Table {table_index} has empty semantic text"
            )

        print("✓ All table results are valid.")

    # ============================================================
    # 7. Replace tables
    # ============================================================

    def insert_paragraph_before_table(self, table, text):

        table_element = table._element

        lines = [
            line.strip()
            for line in text.splitlines()
            if line.strip()
        ]

        for line in lines:

            p = OxmlElement("w:p")
            table_element.addprevious(p)

            paragraph = Paragraph(
                p,
                table._parent,
            )

            paragraph.add_run(line)

        table_element.getparent().remove(
            table_element
        )

    def replace_tables(self, doc, results):

        semantic_text_by_table = {
            result["table_index"]: result["semantic_text"]
            for result in results
        }

        # Reverse order để table index không bị thay đổi
        for table_index in sorted(
            semantic_text_by_table.keys(),
            reverse=True,
        ):

            table = doc.tables[table_index]

            semantic_text = (
                semantic_text_by_table[table_index]
            )

            self.insert_paragraph_before_table(
                table,
                semantic_text,
            )

        print()
        print("✓ All selected tables were replaced with semantic text.")
        print(
            f"Tables processed: "
            f"{len(semantic_text_by_table)}"
        )
        print(
            f"Remaining tables: "
            f"{len(doc.tables)}"
        )

    # ============================================================
    # 8. Process one DOCX
    # ============================================================

    def process(
        self,
        input_path: Path,
        output_path: Path,
    ) -> Path:

        input_path = Path(input_path)
        output_path = Path(output_path)

        print("=" * 80)
        print("DOCX PROCESSING")
        print("=" * 80)

        print(f"Input : {input_path}")
        print(f"Output: {output_path}")

        # Load
        doc = self.load_document(input_path)

        print()
        print(f"Paragraphs: {len(doc.paragraphs)}")
        print(f"Tables    : {len(doc.tables)}")

        # Detect
        tables_need_llm, tables_skipped = (
            self.detect_tables(doc)
        )

        print()
        print(f"Total tables    : {len(doc.tables)}")
        print(
            f"Tables need LLM : "
            f"{len(tables_need_llm)}"
        )
        print(
            f"Tables skipped  : "
            f"{len(tables_skipped)}"
        )

        # Prepare
        tables_for_llm = self.prepare_tables_for_llm(
            doc,
            tables_need_llm,
        )

        print(
            f"Tables prepared for LLM: "
            f"{len(tables_for_llm)}"
        )

        # LLM
        results = self.reconstruct_tables(
            tables_for_llm
        )

        # Validate
        self.validate_results(
            tables_for_llm,
            results,
        )

        # Replace
        self.replace_tables(
            doc,
            results,
        )

        # Save
        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        doc.save(output_path)

        print()
        print(f"✓ Processed DOCX saved to: {output_path}")

        return output_path


# ================================================================
# Manual test
# ================================================================

if __name__ == "__main__":

    INPUT_FILE = Path(
        "input/3352.docx"
    )

    OUTPUT_FILE = Path(
        "output/3352.docx"
    )

    processor = DocxProcessor()

    processor.process(
        input_path=INPUT_FILE,
        output_path=OUTPUT_FILE,
    )

    print()
    print("=" * 80)
    print("TEST PASSED")
    print("=" * 80)